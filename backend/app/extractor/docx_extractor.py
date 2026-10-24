import logging
from pathlib import Path
from app.extractor.base import BaseExtractor, ExtractionResult

logger = logging.getLogger("file_archaeologist.docx_extractor")


class DocxExtractor(BaseExtractor):
    """
    Extracts text paragraphs and table cells from Microsoft Word (.docx) documents.
    """

    def extract(self, file_path: Path) -> ExtractionResult:
        try:
            import docx
            doc = docx.Document(str(file_path))

            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text.strip())

            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        text_parts.append(" | ".join(row_texts))

            full_text = " ".join(text_parts)
            cleaned = self.clean_text(full_text)

            if not cleaned or len(cleaned.strip()) < 10:
                return ExtractionResult(
                    text="", preview="", is_skipped=True,
                    skip_reason="DOCX document contains no text"
                )

            preview = cleaned[:500]
            return ExtractionResult(text=cleaned, preview=preview)

        except Exception as e:
            return ExtractionResult(
                text="", preview="", error=f"Malformed or unreadable DOCX: {str(e)[:200]}"
            )
